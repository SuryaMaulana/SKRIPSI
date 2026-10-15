import json

import numpy as np
import pytest
from openpyxl import load_workbook

from src.data.config import CLASSES
from src.eval import anotasi as an

QUOTA = dict.fromkeys(CLASSES, 1)


def _sample(tmp_path):
    rows = []
    for k, label in enumerate(CLASSES):
        for j in range(2):
            i = f"srbh-{k}{j}"
            rows.append({"id": i, "label": label, "template_key": f"t{k}", "raw_log_masked": f"1.2.3.4 - - [x] \"GET /{i}?q=1\"",
                         "path_norm": f"/{i}", "query_norm": "=cmd|' /C calc'!A0" if k == 0 else "q=1", "ua_norm": "UA"})
    path = tmp_path / "sample.jsonl"
    path.write_text("".join(json.dumps(r) + "\n" for r in rows), encoding="utf-8")
    return path, rows


def _results(tmp_path, rows):
    single, agent = [], []
    for r in rows:
        single.append({"id": r["id"], "label": r["label"], "evidence": [f"/{r['id']}", "GET", 7, ""]})
        agent.append({"id": r["id"], "label": "NORMAL" if r["label"] == "XSS" else r["label"], "evidence": [f"/{r['id']}"]})
    paths = {"llm_single": tmp_path / "llm_single_r1.jsonl", "agent": tmp_path / "agent_r1.jsonl"}
    for method, data in (("llm_single", single), ("agent", agent)):
        paths[method].write_text("".join(json.dumps(x) + "\n" for x in data), encoding="utf-8")
    return paths


@pytest.fixture
def package(tmp_path):
    sample_path, rows = _sample(tmp_path)
    results = _results(tmp_path, rows)
    done = an.prepare(sample_path, tmp_path / "audit", tmp_path / "ids.json", tmp_path / "key.json", results, QUOTA)
    return {"tmp": tmp_path, "sample": sample_path, "rows": {r["id"]: r for r in rows}, "done": done}


def _data(path):
    ws = load_workbook(path)["Anotasi"]
    header = [c.value for c in ws[1]]
    return ws, header, [dict(zip(header, (c.value for c in row))) for row in ws.iter_rows(min_row=2)]


def test_selection_is_stratified_deterministic_and_saved_without_labels(package):
    ids_file = json.loads((package["tmp"] / "ids.json").read_text())
    assert len(ids_file["ids"]) == 6 and ids_file["random_state"] == 42 and "label" not in json.dumps(ids_file["ids"])
    assert sorted(package["rows"][i]["label"] for i in ids_file["ids"]) == sorted(CLASSES)
    again = an.select_ids(an.load_sample(package["sample"]), QUOTA)
    assert again == ids_file["ids"]
    assert sum(an.QUOTA.values()) == 100 and list(an.QUOTA.values()) == [17, 17, 17, 17, 16, 16]


def test_phase_a_files_have_the_requested_columns_no_labels_or_model_output_and_a_dropdown(package):
    files = package["done"]["fase_a"]
    orders = []
    for who, path in files.items():
        ws, header, rows = _data(path)
        assert tuple(header) == an.COLUMNS_A
        assert all(r["label_anotator"] is None and r["catatan"] is None for r in rows)
        values = {str(v) for r in rows for k, v in r.items() if k not in ("label_anotator", "catatan")}
        assert not values & set(CLASSES)  # tidak ada label kelas di data
        assert not any("GET" == v for v in values)  # tidak ada item evidence model
        (validation,) = ws.data_validations.dataValidation
        assert validation.formula1 == '"' + ",".join(CLASSES) + '"' and "G2:G7" in str(validation.sqref)
        formula_cell = [c for c in ws["E"] if isinstance(c.value, str) and c.value.startswith("=cmd")]
        assert formula_cell and formula_cell[0].data_type == "s"  # payload berawalan '=' tetap teks
        guide = [[c.value for c in row] for row in load_workbook(path)["Panduan"].iter_rows()]
        assert [list(r) for r in an.TABLE_1] == [g[:4] for g in guide if g[0] in {r[0] for r in an.TABLE_1}]
        assert all(any(rule in str(g[0]) for g in guide) for rule in an.RULES_A)
        orders.append([r["id"] for r in rows])
    assert sorted(orders[0]) == sorted(orders[1]) and orders[0] != orders[1]


def test_phase_b_hides_methods_merges_identical_items_and_keeps_the_key_separate(package):
    key = json.loads((package["tmp"] / "key.json").read_text())
    items = {x["butir"]: x for x in key["butir"]}
    shared = [x for x in key["butir"] if len(x["asal"]) == 2]
    assert shared and {a["metode"] for a in shared[0]["asal"]} == {"llm_single", "agent"}
    assert all(isinstance(x["evidence_item"], str) for x in key["butir"])  # item non-string dibuang
    for who, path in package["done"]["fase_b"].items():
        ws, header, rows = _data(path)
        assert tuple(header) == an.COLUMNS_B and "metode" not in header
        assert [r["id"] for r in rows] == [items[b]["id"] for b in key["urutan"][who]]
        assert [r["no"] for r in rows] == list(range(1, len(rows) + 1))
        (validation,) = ws.data_validations.dataValidation
        assert validation.formula1 == '"0,1,2"'
    assert key["urutan"]["anotator1"] != key["urutan"]["anotator2"]


def _fill(path, column, values_by_row):
    wb = load_workbook(path)
    ws = wb["Anotasi"]
    header = [c.value for c in ws[1]]
    col = header.index(column) + 1
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=col, value=values_by_row(dict(zip(header, (c.value for c in ws[r])))))
    wb.save(path)


def test_compute_reports_kappa_consensus_and_relevance_from_filled_files(package):
    rows, done, tmp = package["rows"], package["done"], package["tmp"]
    gate = {i: r["label"] for i, r in rows.items()}
    ids = json.loads((tmp / "ids.json").read_text())["ids"]
    disagree, both_wrong = ids[0], ids[1]
    _fill(done["fase_a"]["anotator1"], "label_anotator",
          lambda r: "NORMAL" if r["id"] == both_wrong else gate[r["id"]])
    _fill(done["fase_a"]["anotator2"], "label_anotator",
          lambda r: "NORMAL" if r["id"] in (both_wrong, disagree) else gate[r["id"]].lower())
    _fill(done["fase_a"]["anotator1"], "catatan",
          lambda r: {disagree: "RAGU: SQLI atau NORMAL", both_wrong: "  ragu-ragu"}.get(r["id"]))
    _fill(done["fase_a"]["anotator2"], "catatan", lambda r: "tidak ragu" if r["id"] == disagree else None)
    _fill(done["fase_b"]["anotator1"], "skor", lambda r: 2)
    _fill(done["fase_b"]["anotator2"], "skor", lambda r: "1" if r["evidence_item"] == "GET" else 2)
    files = {"A1": done["fase_a"]["anotator1"], "A2": done["fase_a"]["anotator2"],
             "B1": done["fase_b"]["anotator1"], "B2": done["fase_b"]["anotator2"]}
    a, b, text = an.compute(files, tmp / "ids.json", tmp / "key.json", package["sample"])
    assert a["n"] == 6 and a["persen_setuju"] == pytest.approx(100 * 5 / 6)
    assert a["persen_setuju_anotator1_vs_gerbang"] == pytest.approx(100 * 5 / 6)
    assert a["persen_setuju_anotator2_vs_gerbang"] == pytest.approx(100 * 4 / 6)
    assert [x["id"] for x in a["perlu_diskusi"]] == [disagree]
    assert a["konsensus_beda_gerbang"] == [{"id": both_wrong, "konsensus": "NORMAL", "gerbang": gate[both_wrong]}]
    assert a["kappa_anotator1_vs_gerbang"] == pytest.approx(an.cohen_kappa_score(
        [gate[i] if i != both_wrong else "NORMAL" for i in ids], [gate[i] for i in ids]))
    matrix = a["matriks_konfusi_konsensus_vs_gerbang"]
    assert list(matrix) == list(CLASSES) and list(matrix["SQLI"]) == [*CLASSES, an.DISCUSS]
    assert matrix[gate[disagree]][an.DISCUSS] == 1 and matrix[gate[both_wrong]]["NORMAL"] == 1
    assert sum(matrix[c][c] for c in CLASSES) == 4 and sum(sum(r.values()) for r in matrix.values()) == 6
    assert a["ragu"] == {"anotator1": 2, "anotator2": 0}  # awalan RAGU saja; "tidak ragu" tidak dihitung
    single, agent = b["metode"]["llm_single"], b["metode"]["agent"]
    assert list(b["metode"]) == ["llm_single", "agent"] and b["run"] == {"llm_single": "llm_single_r1",
                                                                         "agent": "agent_r1"}
    assert single["item"] == 18 and agent["item"] == 6  # 6 baris x (path + GET + item kosong) / 6 baris x path
    shown = [r["evidence_item"] for r in _data(done["fase_b"]["anotator1"])[2]]
    assert shown.count(an.EMPTY_ITEM) == 6 and None not in shown
    assert single["rerata_anotator1"] == 2 and single["rerata_anotator2"] == pytest.approx((6 * 1 + 12 * 2) / 18)
    assert single["distribusi_anotator2"] == {0: 0, 1: 6, 2: 12} and agent["rerata"] == 2
    assert agent["item_label_salah"] == 1 and single["item_label_salah"] == 0  # agent memberi NORMAL pada XSS
    per_label = b["metode_per_label"]
    assert list(per_label["llm_single"]) == list(CLASSES) and "XSS" not in per_label["agent"]
    assert per_label["llm_single"]["SQLI"]["distribusi_anotator2"] == {0: 0, 1: 1, 2: 2}
    assert per_label["llm_single"]["SQLI"]["rerata_anotator2"] == pytest.approx(5 / 3)
    assert per_label["agent"]["NORMAL"]["item"] == 2 and per_label["agent"]["NORMAL"]["item_label_salah"] == 1
    assert sum(x["item"] for x in per_label["llm_single"].values()) == single["item"]
    assert "Cohen's κ antar anotator" in text and "κ berbobot kuadratik" in text and "Kesepakatan mentah" in text
    assert "Matriks konfusi" in text and "catatan RAGU" in text and "anotator 1 2, anotator 2 0" in text
    assert "| LLM tunggal | SQLI | 3 | 1 | 1.83 | 2.00 | 1.67 | 0/0/3 | 0/1/2 |" in text
    assert "LLM tunggal = llm_single_r1; deep agent = agent_r1" in text
    # 19 butir: 6 baris x 3 item LLM tunggal + 1 item agent (label NORMAL pada baris XSS); 6 butir "GET" berselisih 1
    assert b["kesepakatan_mentah"] == {"setuju": 13, "dari": 19, "persen": pytest.approx(100 * 13 / 19), "selisih_1": 6,
                                       "selisih_2": 0, "tabel_silang": {0: {0: 0, 1: 0, 2: 0}, 1: {0: 0, 1: 0, 2: 0},
                                                                        2: {0: 0, 1: 6, 2: 13}}}
    assert ("Kesepakatan mentah antar anotator: 68.4% (13 dari 19 butir); dari 6 butir yang berbeda, 6 berselisih 1 "
            "dan 0 berselisih 2.") in text and "| 2 | 0 | 6 | 13 | 19 |" in text
    diff = b["selisih_metode"]  # tiap baris berkomposisi sama, jadi semua draw bernilai sama dengan selisih teramati
    assert diff["klaster"] == diff["baris_berpasangan"] == 6 and diff["n_draw"] == an.N_DRAW
    assert diff["arah"] == "deep agent - LLM tunggal"
    item, row = diff["per_item"], diff["per_baris"]
    assert item["rata_rata_dua_anotator"]["selisih"] == pytest.approx(agent["rerata"] - single["rerata"])
    assert item["anotator2"]["selisih"] == pytest.approx(agent["rerata_anotator2"] - single["rerata_anotator2"])
    assert item["anotator1"] == {"selisih": 0, "ik_bawah": 0, "ik_atas": 0, "p": 1.0}
    assert item["anotator2"]["ik_bawah"] == pytest.approx(item["anotator2"]["ik_atas"]) and item["anotator2"]["p"] == 0
    assert all(row[k] == pytest.approx(item[k]) for k in item)  # komposisi sama per baris: kedua bobot sama
    assert "| rata-rata dua anotator | 0.167 | [0.167; 0.167] | <0.001 | 0.167 | [0.167; 0.167] | <0.001 |" in text
    assert "| anotator 1 | 0.000 | [0.000; 0.000] | 1.000 | 0.000 | [0.000; 0.000] | 1.000 |" in text
    # label gerbang per id hanya di results/: laporan reports/ memuat agregat dan daftar "perlu diskusi" saja
    assert both_wrong not in text and disagree in text and "results/revisi_anotasi_per_baris.json" in text
    an.write_outputs(a, b, text, tmp / "reports" / "anotasi.md", tmp / "results" / "per_baris.json")
    public = json.loads((tmp / "reports" / "anotasi.json").read_text())
    assert "konsensus" not in public["fase_a"] and "konsensus_beda_gerbang" not in public["fase_a"]
    assert both_wrong not in json.dumps(public) and public["fase_a"]["jumlah_konsensus_beda_gerbang"] == 1
    assert public["fase_a"]["konsensus_beda_gerbang_per_kelas"] == {f"{gate[both_wrong]} -> NORMAL": 1}
    per_row = json.loads((tmp / "results" / "per_baris.json").read_text())
    assert per_row["konsensus_beda_gerbang"] == a["konsensus_beda_gerbang"] and len(per_row["konsensus"]) == 6
    assert "asal" not in json.dumps(per_row) and "metode" not in json.dumps(per_row)  # kunci Fase B tidak ikut


def _key_for_bootstrap():
    """5 baris dengan jumlah item berbeda; butir B0 milik kedua metode, B5 muncul dua kali di keluaran agent."""
    asal = {"L": [{"metode": "llm_single"}], "A": [{"metode": "agent"}],
            "LA": [{"metode": "llm_single"}, {"metode": "agent"}], "AA": [{"metode": "agent"}, {"metode": "agent"}]}
    spec = [("r0", "LA"), ("r0", "L"), ("r1", "A"), ("r1", "L"), ("r1", "L"), ("r2", "AA"), ("r2", "L"), ("r3", "A"),
            ("r3", "L"), ("r4", "LA"), ("r4", "A"), ("r4", "L")]
    butir = [{"butir": f"B{k}", "id": i, "asal": asal[a]} for k, (i, a) in enumerate(spec)]
    s1 = dict(zip((x["butir"] for x in butir), (2, 1, 2, 0, 1, 2, 1, 2, 2, 1, 2, 0)))
    s2 = dict(zip((x["butir"] for x in butir), (2, 2, 1, 0, 1, 2, 0, 2, 1, 1, 2, 1)))
    return {"butir": butir}, s1, s2


def test_method_difference_is_a_cluster_bootstrap_over_log_rows():
    key, s1, s2 = _key_for_bootstrap()
    ids = ["r3", "r1", "r4", "r0", "r2"]
    out = an.method_difference(key, s1, s2, ids, n=300)
    scorers = {"rata_rata_dua_anotator": lambda b: (s1[b] + s2[b]) / 2, "anotator1": s1.get, "anotator2": s2.get}

    def mean(rows, method, score):
        values = [score(x["butir"]) for r in rows for x in key["butir"] if x["id"] == r for a in x["asal"]
                  if a["metode"] == method]
        return sum(values) / len(values)

    def per_item(rows, score):  # rerata atas item dari baris terambil (baris ganda ikut berulang)
        return mean(rows, "agent", score) - mean(rows, "llm_single", score)

    def per_row(rows, score):  # rerata atas baris terambil dari selisih rerata kedua metode pada baris itu
        return sum(per_item([r], score) for r in rows) / len(rows)

    rows = sorted(ids)  # klaster diurut menurut id, lalu satu rng.integers atas semua baris per draw (tanpa strata)
    rng = np.random.default_rng(an.RANDOM_STATE)
    draws = [[rows[k] for k in rng.integers(0, len(rows), len(rows))] for _ in range(300)]
    for name, score in scorers.items():
        for weighting, estimate in (("per_item", per_item), ("per_baris", per_row)):
            expected = np.array([estimate(d, score) for d in draws])
            expected[np.abs(expected) < 1e-9] = 0.0  # seperti modul: nol aritmetika tidak bergantung pembulatan
            low, high = np.percentile(expected, [2.5, 97.5])
            p = min(1.0, 2 * min(np.mean(expected <= 0), np.mean(expected >= 0)))
            got = out[weighting][name]
            assert got["selisih"] == pytest.approx(estimate(rows, score))
            assert (got["ik_bawah"], got["ik_atas"], got["p"]) == pytest.approx((low, high, p))
    assert out["per_item"]["anotator1"]["selisih"] != pytest.approx(out["per_baris"]["anotator1"]["selisih"])
    assert out["klaster"] == out["baris_berpasangan"] == 5 and out["n_draw"] == 300
    assert out == an.method_difference(key, s1, s2, list(reversed(ids)), n=300)  # urutan masukan tidak berpengaruh
    same = an.method_difference(key, s1, dict(s1), ids, n=300)  # draw yang sama untuk ketiga skor
    for weighting in ("per_item", "per_baris"):
        assert same[weighting]["anotator1"] == same[weighting]["anotator2"] == same[weighting]["rata_rata_dua_anotator"]
    only_llm = {"butir": [x for x in key["butir"] if all(a["metode"] == "llm_single" for a in x["asal"])]}
    assert an.method_difference(only_llm, s1, s2, ids) is None


def test_score_agreement_counts_absolute_gaps_in_both_directions():
    s1 = {"a": 2, "b": 0, "c": 1, "d": 2, "e": 0, "f": 1}
    s2 = {"a": 0, "b": 2, "c": 1, "d": 1, "e": 1, "f": 1}
    out = an.score_agreement(s1, s2)
    assert (out["setuju"], out["dari"], out["selisih_1"], out["selisih_2"]) == (2, 6, 2, 2)
    assert out["persen"] == pytest.approx(100 * 2 / 6)
    assert out["tabel_silang"] == {0: {0: 0, 1: 1, 2: 1}, 1: {0: 0, 1: 2, 2: 0}, 2: {0: 1, 1: 1, 2: 0}}


def test_prepare_key_writes_the_same_ids_and_key_as_prepare(package, monkeypatch):
    tmp = package["tmp"]
    results = {"llm_single": tmp / "llm_single_r1.jsonl", "agent": tmp / "agent_r1.jsonl"}

    def no_sheets(*args, **kwargs):
        raise AssertionError("prepare_key tidak boleh membuat berkas anotasi")

    monkeypatch.setattr(an, "build_phase_a", no_sheets)
    monkeypatch.setattr(an, "build_phase_b", no_sheets)
    done = an.prepare_key(package["sample"], tmp / "ids2.json", tmp / "key2.json", results, QUOTA)
    assert (tmp / "ids2.json").read_bytes() == (tmp / "ids.json").read_bytes()
    assert (tmp / "key2.json").read_bytes() == (tmp / "key.json").read_bytes()
    assert done["butir"] == package["done"]["butir"]


def test_compute_refuses_incomplete_or_edited_files(package):
    done, tmp = package["done"], package["tmp"]
    files = {"A1": done["fase_a"]["anotator1"], "A2": done["fase_a"]["anotator2"],
             "B1": done["fase_b"]["anotator1"], "B2": done["fase_b"]["anotator2"]}
    with pytest.raises(ValueError, match="label kosong"):
        an.compute(files, tmp / "ids.json", tmp / "key.json", package["sample"])
    for who in ("anotator1", "anotator2"):
        _fill(done["fase_a"][who], "label_anotator", lambda r: "SQLI")
        _fill(done["fase_b"][who], "skor", lambda r: 2)
    _fill(done["fase_b"]["anotator2"], "evidence_item", lambda r: r["evidence_item"] + " (diubah)")
    with pytest.raises(ValueError, match="tidak cocok dengan kunci"):
        an.compute(files, tmp / "ids.json", tmp / "key.json", package["sample"])
