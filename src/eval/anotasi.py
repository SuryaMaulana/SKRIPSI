"""Paket audit manusia untuk revisi jurnal (keputusan penulis 2026-10-13). Pelabelan dan penilaian dikerjakan dua anotator;
skrip ini hanya menyiapkan berkas dan menghitung hasilnya.

siapkan:
- 100 baris test_sample_500 dipilih stratified per kelas (SQLI 17, XSS 17, PATH_TRAVERSAL 17, CMD_INJECTION 17,
  WEB_SCANNING 16, NORMAL 16): baris tiap kelas diurut menurut id lalu diambil acak dengan random_state 42; daftar id
  -> data/processed/audit_100_ids.json (tanpa label).
- Fase A (label): audit/anotasi_A_anotator1.xlsx dan audit/anotasi_A_anotator2.xlsx berisi baris yang sama dengan urutan
  acak berbeda per anotator (seed 1 dan 2). Kolom: no, id, raw_log_masked, path_norm, query_norm, ua_norm,
  label_anotator (dropdown 6 kelas), catatan. Tidak memuat label kelas maupun keluaran model. Lembar Panduan berisi
  Tabel 1 manuskrip apa adanya dan aturan pelabelan penulis.
- Fase B (relevansi bukti): item evidence llm_single_r1 dan agent_r1 untuk 100 baris yang sama, satu baris per item
  unik (id, label_prediksi, evidence_item); item yang identik di kedua metode digabung. Metode disamarkan: tidak ada
  kolom metode, urutan acak berbeda per anotator (seed 11 dan 12). Kolom: no, id, raw_log_masked, label_prediksi,
  evidence_item, skor (dropdown 0/1/2), catatan. Pemetaan butir -> metode disimpan terpisah di
  data/processed/audit_B_key.json, yang tidak dibagikan ke anotator.
hitung: membaca keempat berkas terisi -> reports/revisi_anotasi.md (+ .json) dan results/revisi_anotasi_per_baris.json:
- Fase A: kappa Cohen tak berbobot antar anotator dan kesepakatan mentah (persen setuju); kappa dan kesepakatan mentah
  tiap anotator vs label gerbang bukti; daftar baris "perlu diskusi" (id dan label kedua anotator); jumlah baris yang
  konsensusnya berbeda dari label gerbang per pasangan kelas; matriks konfusi konsensus vs label gerbang per kelas
  (agregat); jumlah baris dengan catatan berawalan RAGU per anotator. Label konsensus per id dan label gerbang per id
  pada baris yang berbeda hanya ditulis ke results/revisi_anotasi_per_baris.json (keputusan penulis: label gerbang
  sampel uji tidak ditulis di luar results/).
- Fase B: rerata skor relevansi per metode setelah kunci dibuka, kappa berbobot kuadratik dan kesepakatan mentah antar
  anotator (berikut tabel silang skor dan jumlah butir yang berselisih 1 atau 2), distribusi skor 0/1/2 per metode,
  rerata skor pada baris yang labelnya benar vs salah, rerata serta distribusi per metode dan label prediksi, dan
  selisih rerata skor deep agent - LLM tunggal, per item (rerata atas item seperti tabel per metode) dan per baris log
  (rerata atas baris dari selisih rerata kedua metode pada baris yang sama, setiap baris berbobot sama), dengan IK 95%
  bootstrap klaster per baris log: baris audit diurut menurut id, 10.000 draw dengan default_rng(42), tiap draw
  mengambil baris sebanyak jumlah baris dengan pengembalian, dan indeks yang sama dipakai untuk kedua pembobotan, skor
  rata-rata dua anotator, dan skor tiap anotator; IK = persentil 2,5/97,5; p = 2 x porsi terkecil draw dengan selisih
  <= 0 atau >= 0 (seperti src/eval/revisi.py).
kunci: daftar id dan kunci Fase B saja, tanpa berkas anotasi, identik dengan keluaran siapkan; make eval memanggilnya
  bila data/processed/audit_*.json belum ada, lalu menghitung dari lembar terisi di audit/terisi/.
Jalankan: python -m src.eval.anotasi siapkan | kunci | hitung --a1 ... --a2 ... --b1 ... --b2 ...  (atau make eval)
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import statistics
import sys
from collections import Counter
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.metrics import cohen_kappa_score

from ..common.inputs import SAMPLE_PATH
from ..data.config import CLASSES, RANDOM_STATE, ROOT

QUOTA = {"SQLI": 17, "XSS": 17, "PATH_TRAVERSAL": 17, "CMD_INJECTION": 17, "WEB_SCANNING": 16, "NORMAL": 16}
IDS_PATH = ROOT / "data" / "processed" / "audit_100_ids.json"
KEY_PATH = ROOT / "data" / "processed" / "audit_B_key.json"
PER_ROW_PATH = ROOT / "results" / "revisi_anotasi_per_baris.json"
OUT_DIR = ROOT / "audit"
RESULTS = ROOT / "results"
METHOD_RUNS = {"llm_single": "llm_single_r1", "agent": "agent_r1"}
SEEDS_A = {"anotator1": 1, "anotator2": 2}
SEEDS_B = {"anotator1": 11, "anotator2": 12}
N_DRAW = 10_000
COLUMNS_A = ("no", "id", "raw_log_masked", "path_norm", "query_norm", "ua_norm", "label_anotator", "catatan")
COLUMNS_B = ("no", "id", "raw_log_masked", "label_prediksi", "evidence_item", "skor", "catatan")
SCORES = (0, 1, 2)
DISCUSS = "perlu diskusi"
EMPTY_ITEM = "[item kosong]"  # item evidence berupa string kosong/spasi; sel kosong di xlsx terbaca sebagai None
METHOD_NAMES = {"llm_single": "LLM tunggal", "agent": "deep agent"}
RAGU = re.compile(r"\s*RAGU\b", re.IGNORECASE)  # audit/PANDUAN_ANOTASI.md: catatan baris yang ragu diawali "RAGU:"

# Tabel 1 manuskrip revisi (revisi_v2.docx), disalin apa adanya.
TABLE_1 = (
    ("Class", "Code", "Definition", "Indicators in request, user agent, or referer"),
    ("SQL injection", "SQLI", "SQL syntax injected into a parameter to manipulate a database query",
     "' OR 1=1, UNION SELECT, SLEEP(, comments -- or #, often URL-encoded (%27, %20)"),
    ("Cross-site scripting", "XSS", "Client-side script injected for execution in the victim's browser",
     "<script>, onerror=, javascript:, encoded variants (%3Cscript%3E), may appear in the referer"),
    ("Path traversal / file inclusion", "PATH_TRAVERSAL",
     "Access to files outside the web root, or local/remote file inclusion",
     "../, ..%2f, /etc/passwd, win.ini, external URLs with a scheme (http://...) in a path or parameter (RFI)"),
    ("Command injection", "CMD_INJECTION", "OS commands injected through a parameter",
     "Command separators (;, &&, |) followed by shell commands such as id, cat /etc/passwd, sleep 15; backtick or $(...) "
     "substitution; encoded variants (%7C, %26%26, %60); Shellshock; PHP-CGI -d allow_url_include"),
    ("Web scanning / enumeration", "WEB_SCANNING", "Automated probing for files, directories, or server capabilities",
     "Scanner user agents (Nmap, zgrab, sqlmap), non-standard methods (PROPFIND, MKCOL, TRACK), backup/VCS files (.bak, "
     ".swp, .git/), probes for /.env, /wp-login.php"),
    ("Normal", "NORMAL", "Legitimate user or search-engine crawler request",
     "Ordinary URI, alphanumeric parameters, status 200/301/304, browser or verified bot user agent"),
)
RULES_A = (
    "Beri label hanya dari apa yang terlihat di baris.",
    "Bila ragu antara serangan dan NORMAL, pilih NORMAL.",
    "Bila ragu antara dua kelas serangan, pilih yang payload-nya paling eksplisit.",
)
HOWTO_A = (
    "Pilih satu kelas di kolom label_anotator untuk setiap baris (dropdown). Kolom catatan boleh dikosongkan.",
    "Jangan mengubah kolom lain dan jangan menghapus baris; mengurutkan ulang baris boleh.",
    "Kerjakan sendiri tanpa berdiskusi dengan anotator lain dan tanpa melihat label dataset atau keluaran model sampai "
    "kedua berkas selesai.",
)
SCALE_B = (
    ("0", "Tidak ada di baris log, atau tidak relevan."),
    ("1", "Ada di baris log, tetapi tidak menunjukkan kelas yang diprediksi."),
    ("2", "Ada di baris log dan menunjukkan kelas yang diprediksi."),
)
HOWTO_B = (
    "Nilai setiap item evidence terhadap label_prediksi di baris yang sama, bukan terhadap label yang menurut Anda benar.",
    "Item dianggap ada bila muncul di baris log apa adanya atau setelah URL-decode (mis. ' OR 1=1 untuk %27%20OR%201=1).",
    "Pilih skor 0, 1, atau 2 di kolom skor (dropdown). Kolom catatan boleh dikosongkan.",
    f"{EMPTY_ITEM} berarti model menulis item evidence kosong (tanpa isi).",
    "Jangan mengubah kolom lain dan jangan menghapus baris; mengurutkan ulang baris boleh. Kerjakan sendiri.",
)
# Dua contoh netral: baris rekaan (alamat IP dokumentasi RFC 5737), bukan dari dataset maupun keluaran metode.
EXAMPLES_B = (
    ('203.0.113.5 - - [01/Jan/2020:00:00:00 +0000] "GET /produk.php?id=7%20UNION%20SELECT%20username,password%20FROM'
     '%20users HTTP/1.1" 200 512 "-" "Mozilla/5.0"', "SQLI",
     (("UNION SELECT username,password", 2, "ada setelah URL-decode dan menunjukkan SQL injection"),
      ("GET /produk.php", 1, "ada, tetapi tidak menunjukkan SQL injection"),
      ("DROP TABLE users", 0, "tidak ada di baris log"))),
    ('198.51.100.7 - - [01/Jan/2020:00:00:00 +0000] "GET /blog/index.php/2020/03/22/hello-world/ HTTP/1.1" 200 8042 "-" '
     '"Mozilla/5.0 (Windows NT 10.0; Win64; x64)"', "NORMAL",
     (("/blog/index.php/2020/03/22/hello-world/", 2, "ada dan menunjukkan permintaan halaman yang wajar"),
      ("GET", 1, "ada, tetapi tidak membedakan permintaan normal dari serangan"),
      ("<script>", 0, "tidak ada di baris log"))),
)

HEADER_FONT = Font(bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
WRAP = Alignment(wrap_text=True, vertical="top")


# ------------------------------------------------------------------------------------------------- siapkan


def _clean(value):
    """Teks aman untuk xlsx: karakter kontrol yang tidak sah di XML ditulis sebagai \\xhh."""
    if isinstance(value, str):
        return ILLEGAL_CHARACTERS_RE.sub(lambda m: f"\\x{ord(m.group()):02x}", value)
    return value


def _put(ws, row: int, col: int, value):
    cell = ws.cell(row=row, column=col, value=_clean(value))
    if isinstance(value, str):
        cell.data_type = "s"  # jangan pernah ditafsirkan sebagai rumus (teks payload bisa diawali '=')
    return cell


def load_sample(path: Path = SAMPLE_PATH) -> pd.DataFrame:
    return pd.read_json(path, lines=True, dtype={"id": str})


def select_ids(sample: pd.DataFrame, quota: dict[str, int] = QUOTA, random_state: int = RANDOM_STATE) -> list[str]:
    ids = []
    for label in CLASSES:
        rows = sample[sample["label"] == label].sort_values("id")
        if len(rows) < quota[label]:
            raise ValueError(f"{label}: hanya {len(rows)} baris, kuota {quota[label]}")
        ids += rows.sample(n=quota[label], random_state=random_state)["id"].tolist()
    return ids


def _shuffled(items: list, seed: int) -> list:
    order = np.random.default_rng(seed).permutation(len(items))
    return [items[k] for k in order]


def _sheet(wb: Workbook, title: str, columns: tuple[str, ...], rows: list[list], dropdown: tuple[str, list[str]],
           widths: dict[str, int]):
    ws = wb.active
    ws.title = title
    for c, name in enumerate(columns, 1):
        _put(ws, 1, c, name).font = HEADER_FONT
    for r, values in enumerate(rows, 2):
        for c, value in enumerate(values, 1):
            _put(ws, r, c, value).alignment = WRAP
    column, options = dropdown
    letter = ws.cell(row=1, column=columns.index(column) + 1).column_letter
    validation = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True,
                                showErrorMessage=True, errorTitle="Nilai tidak sah",
                                error="Pilih salah satu nilai dari daftar.")
    ws.add_data_validation(validation)
    if rows:
        validation.add(f"{letter}2:{letter}{len(rows) + 1}")
    for name in (column, "catatan"):
        idx = columns.index(name) + 1
        for r in range(2, len(rows) + 2):
            ws.cell(row=r, column=idx).fill = INPUT_FILL
    for c, name in enumerate(columns, 1):
        ws.column_dimensions[ws.cell(row=1, column=c).column_letter].width = widths.get(name, 14)
    ws.freeze_panes = "A2"
    return ws


def _guide(wb: Workbook, lines: list[list[str]]):
    ws = wb.create_sheet("Panduan")
    for r, values in enumerate(lines, 1):
        for c, value in enumerate(values, 1):
            _put(ws, r, c, value).alignment = WRAP
    for letter, width in zip("ABCD", (34, 18, 50, 80)):
        ws.column_dimensions[letter].width = width
    return ws


def build_phase_a(sample: pd.DataFrame, ids: list[str], seed: int, path: Path) -> list[str]:
    by_id = sample.set_index("id")
    order = _shuffled(ids, seed)
    rows = [[k, i, *(by_id.at[i, c] if pd.notna(by_id.at[i, c]) else "" for c in COLUMNS_A[2:6]), "", ""]
            for k, i in enumerate(order, 1)]
    wb = Workbook()
    _sheet(wb, "Anotasi", COLUMNS_A, rows, ("label_anotator", list(CLASSES)),
           {"no": 6, "id": 24, "raw_log_masked": 70, "path_norm": 40, "query_norm": 40, "ua_norm": 40,
            "label_anotator": 20, "catatan": 30})
    guide = [["FASE A — LABEL KELAS"], [""], ["Definisi kelas (Tabel 1 manuskrip)"], *[list(r) for r in TABLE_1], [""],
             ["Aturan pelabelan"], *[[f"{k}. {rule}"] for k, rule in enumerate(RULES_A, 1)], [""], ["Cara mengisi"],
             *[[f"- {line}"] for line in HOWTO_A]]
    _guide(wb, guide)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return order


def evidence_items(results: dict[str, Path], ids: list[str]) -> list[dict]:
    """Butir unik (id, label_prediksi, evidence_item) beserta asal metodenya, urut id lalu kemunculan."""
    wanted, items, index = set(ids), [], {}
    for method, path in results.items():
        for line in path.read_text(encoding="utf-8").splitlines():
            row = json.loads(line)
            if row["id"] not in wanted:
                continue
            for k, item in enumerate(e for e in row.get("evidence") or [] if isinstance(e, str)):
                key = (row["id"], row["label"], item)
                if key not in index:
                    index[key] = len(items)
                    items.append({"id": row["id"], "label_prediksi": row["label"], "evidence_item": item, "asal": []})
                items[index[key]]["asal"].append({"metode": method, "run": path.stem, "indeks_item": k})
    position = {i: k for k, i in enumerate(ids)}
    items.sort(key=lambda x: position[x["id"]])
    for k, item in enumerate(items, 1):
        item["butir"] = f"B{k:04d}"
    return items


def shown_item(item: str) -> str:
    """Teks item evidence di berkas anotator (item kosong diberi penanda agar tidak hilang menjadi sel kosong)."""
    return _clean(item) if item.strip() else EMPTY_ITEM


def build_phase_b(sample: pd.DataFrame, items: list[dict], seed: int, path: Path) -> list[str]:
    logs = dict(zip(sample["id"], sample["raw_log_masked"]))
    order = _shuffled(items, seed)
    rows = [[k, x["id"], logs[x["id"]], x["label_prediksi"], shown_item(x["evidence_item"]), "", ""]
            for k, x in enumerate(order, 1)]
    wb = Workbook()
    _sheet(wb, "Anotasi", COLUMNS_B, rows, ("skor", [str(s) for s in SCORES]),
           {"no": 6, "id": 24, "raw_log_masked": 70, "label_prediksi": 18, "evidence_item": 40, "skor": 8, "catatan": 30})
    guide = [["FASE B — RELEVANSI BUKTI"], [""], ["Skala skor"], *[[s, d] for s, d in SCALE_B], [""], ["Cara menilai"],
             *[[f"- {line}"] for line in HOWTO_B], [""], ["Contoh netral (baris rekaan, bukan dari data audit)"]]
    for k, (line, label, cases) in enumerate(EXAMPLES_B, 1):
        guide += [[f"Contoh {k}", f"label_prediksi: {label}"], ["baris log", line],
                  *[[f'item "{item}"', f"skor {score}", reason] for item, score, reason in cases], [""]]
    guide += [["Definisi kelas (Tabel 1 manuskrip)"], *[list(r) for r in TABLE_1]]
    _guide(wb, guide)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return [x["butir"] for x in order]


def _save_ids(ids: list[str], quota: dict[str, int], sample_path: Path, ids_path: Path) -> None:
    source_sha = hashlib.sha256(sample_path.read_bytes()).hexdigest()
    ids_path.parent.mkdir(parents=True, exist_ok=True)
    ids_path.write_text(json.dumps({"ids": ids, "kuota": quota, "random_state": RANDOM_STATE,
                                    "sumber": str(sample_path.relative_to(ROOT)) if sample_path.is_relative_to(ROOT)
                                    else str(sample_path), "sumber_sha256": source_sha}, indent=1) + "\n",
                        encoding="utf-8")


def _save_key(items: list[dict], orders: dict[str, list[str]], results: dict[str, Path], key_path: Path) -> None:
    key = {"keterangan": "Kunci Fase B: jangan dibagikan ke anotator. 'urutan' memberi butir untuk setiap nomor baris "
                         "(no = posisi + 1) di berkas anotator.",
           "seed_fase_a": SEEDS_A, "seed_fase_b": SEEDS_B, "run": {m: p.stem for m, p in results.items()},
           "butir": items, "urutan": orders}
    key_path.parent.mkdir(parents=True, exist_ok=True)
    key_path.write_text(json.dumps(key, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def prepare(sample_path: Path = SAMPLE_PATH, out_dir: Path = OUT_DIR, ids_path: Path = IDS_PATH,
            key_path: Path = KEY_PATH, results: dict[str, Path] | None = None, quota: dict[str, int] = QUOTA) -> dict:
    sample = load_sample(sample_path)
    ids = select_ids(sample, quota)
    _save_ids(ids, quota, sample_path, ids_path)
    files_a = {who: out_dir / f"anotasi_A_{who}.xlsx" for who in SEEDS_A}
    for who, path in files_a.items():
        build_phase_a(sample, ids, SEEDS_A[who], path)
    results = results or {m: RESULTS / f"{run}.jsonl" for m, run in METHOD_RUNS.items()}
    items = evidence_items(results, ids)
    files_b = {who: out_dir / f"anotasi_B_{who}.xlsx" for who in SEEDS_B}
    orders = {who: build_phase_b(sample, items, SEEDS_B[who], path) for who, path in files_b.items()}
    _save_key(items, orders, results, key_path)
    return {"ids": ids, "fase_a": files_a, "fase_b": files_b, "kunci": key_path, "butir": len(items)}


def prepare_key(sample_path: Path = SAMPLE_PATH, ids_path: Path = IDS_PATH, key_path: Path = KEY_PATH,
                results: dict[str, Path] | None = None, quota: dict[str, int] = QUOTA) -> dict:
    """Seperti prepare, tetapi hanya daftar id dan kunci (tanpa berkas anotasi); urutan butir dari _shuffled yang sama
    dengan build_phase_b, sehingga kedua berkas JSON identik dengan keluaran prepare."""
    sample = load_sample(sample_path)
    ids = select_ids(sample, quota)
    _save_ids(ids, quota, sample_path, ids_path)
    results = results or {m: RESULTS / f"{run}.jsonl" for m, run in METHOD_RUNS.items()}
    items = evidence_items(results, ids)
    _save_key(items, {who: [x["butir"] for x in _shuffled(items, seed)] for who, seed in SEEDS_B.items()}, results,
              key_path)
    return {"ids": ids, "kunci": key_path, "butir": len(items)}


# ------------------------------------------------------------------------------------------------- hitung


def read_sheet(path: Path, columns: tuple[str, ...]) -> list[dict]:
    ws = load_workbook(path, data_only=True)["Anotasi"]
    header = [str(c.value).strip() if c.value is not None else "" for c in ws[1]]
    missing = [c for c in columns if c not in header]
    if missing:
        raise ValueError(f"{path.name}: kolom hilang {missing}")
    pos = {c: header.index(c) for c in columns}
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if all(v is None or str(v).strip() == "" for v in values):
            continue
        rows.append({c: values[pos[c]] for c in columns})
    return rows


def _label(value) -> str | None:
    text = re.sub(r"\s+", "", str(value or "")).upper()
    return text if text in CLASSES else None


def _score(value) -> int | None:
    try:
        number = float(str(value).strip())
    except ValueError:
        return None
    return int(number) if number in SCORES else None


def labels_a(path: Path, ids: list[str]) -> dict[str, str]:
    rows = read_sheet(path, COLUMNS_A)
    found = Counter(str(r["id"]) for r in rows)
    problems = [f"id ganda {i}" for i, n in found.items() if n > 1]
    problems += [f"id tidak dikenal {i}" for i in found if i not in set(ids)]
    problems += [f"id hilang {i}" for i in ids if i not in found]
    out = {}
    for r in rows:
        label = _label(r["label_anotator"])
        if label is None:
            problems.append(f"no {r['no']} ({r['id']}): label kosong/tidak sah {r['label_anotator']!r}")
        out[str(r["id"])] = label
    if problems:
        raise ValueError(f"{path.name}: " + "; ".join(problems[:20]) + (" ..." if len(problems) > 20 else ""))
    return out


def scores_b(path: Path, key: dict, who: str) -> dict[str, int]:
    rows = read_sheet(path, COLUMNS_B)
    items = {x["butir"]: x for x in key["butir"]}
    order = key["urutan"][who]
    out, problems = {}, []
    for r in rows:
        try:
            butir = order[int(r["no"]) - 1]
        except (TypeError, ValueError, IndexError):
            problems.append(f"no tidak sah {r['no']!r}")
            continue
        x = items[butir]
        if (str(r["id"]), str(r["label_prediksi"]), str(r["evidence_item"])) != \
                (x["id"], x["label_prediksi"], shown_item(x["evidence_item"])):
            problems.append(f"no {r['no']}: isi baris tidak cocok dengan kunci")
            continue
        score = _score(r["skor"])
        if score is None:
            problems.append(f"no {r['no']}: skor kosong/tidak sah {r['skor']!r}")
        if butir in out:
            problems.append(f"no {r['no']}: baris ganda")
        out[butir] = score
    problems += [f"butir hilang {b}" for b in order if b not in out]
    if problems:
        raise ValueError(f"{path.name}: " + "; ".join(problems[:20]) + (" ..." if len(problems) > 20 else ""))
    return out


def ragu_rows(path: Path) -> list[str]:
    """id baris Fase A yang catatannya diawali RAGU (spasi di depan dan huruf besar/kecil diabaikan)."""
    return [str(r["id"]) for r in read_sheet(path, COLUMNS_A) if RAGU.match(str(r["catatan"] or ""))]


def _agreement(x: list[str], y: list[str]) -> float:
    return 100 * float(np.mean([a == b for a, b in zip(x, y)]))


def confusion(consensus: dict[str, str], gate: dict[str, str], ids: list[str]) -> dict[str, dict[str, int]]:
    """Baris = label gerbang, kolom = konsensus (6 kelas + "perlu diskusi"); hanya jumlah, tanpa id."""
    matrix = {g: dict.fromkeys([*CLASSES, DISCUSS], 0) for g in CLASSES}
    for i in ids:
        matrix[gate[i]][consensus[i]] += 1
    return matrix


def phase_a_report(first: dict[str, str], second: dict[str, str], gate: dict[str, str], ids: list[str],
                   ragu: dict[str, list[str]] | None = None) -> dict:
    l1, l2, g = ([d[i] for i in ids] for d in (first, second, gate))
    consensus = {i: (a if a == b else DISCUSS) for i, a, b in zip(ids, l1, l2)}
    return {
        "n": len(ids),
        "kappa_antar_anotator": float(cohen_kappa_score(l1, l2)),
        "persen_setuju": _agreement(l1, l2),
        "kappa_anotator1_vs_gerbang": float(cohen_kappa_score(l1, g)),
        "kappa_anotator2_vs_gerbang": float(cohen_kappa_score(l2, g)),
        "persen_setuju_anotator1_vs_gerbang": _agreement(l1, g),
        "persen_setuju_anotator2_vs_gerbang": _agreement(l2, g),
        "perlu_diskusi": [{"id": i, "anotator1": first[i], "anotator2": second[i]} for i in ids
                          if consensus[i] == DISCUSS],
        "konsensus_beda_gerbang": [{"id": i, "konsensus": consensus[i], "gerbang": gate[i]} for i in ids
                                   if consensus[i] not in (DISCUSS, gate[i])],
        "matriks_konfusi_konsensus_vs_gerbang": confusion(consensus, gate, ids),
        "ragu": {who: len(rows) for who, rows in (ragu or {}).items()},
        "konsensus": consensus,
    }


def _relevance(inst: list[tuple[dict, dict]], s1: dict[str, int], s2: dict[str, int], gate: dict[str, str]) -> dict:
    """Statistik skor atas instans (butir, asal); butir milik kedua metode ikut di instans kedua metode."""
    def both(x: dict) -> float:
        return (s1[x["butir"]] + s2[x["butir"]]) / 2

    correct = [both(x) for x, _ in inst if x["label_prediksi"] == gate[x["id"]]]
    wrong = [both(x) for x, _ in inst if x["label_prediksi"] != gate[x["id"]]]
    return {
        "item": len(inst), "baris": len({x["id"] for x, _ in inst}),
        "rerata_anotator1": statistics.mean(s1[x["butir"]] for x, _ in inst),
        "rerata_anotator2": statistics.mean(s2[x["butir"]] for x, _ in inst),
        "rerata": statistics.mean([both(x) for x, _ in inst]),
        "distribusi_anotator1": {s: sum(s1[x["butir"]] == s for x, _ in inst) for s in SCORES},
        "distribusi_anotator2": {s: sum(s2[x["butir"]] == s for x, _ in inst) for s in SCORES},
        "rerata_label_benar": statistics.mean(correct) if correct else None, "item_label_benar": len(correct),
        "rerata_label_salah": statistics.mean(wrong) if wrong else None, "item_label_salah": len(wrong),
    }


def score_agreement(s1: dict[str, int], s2: dict[str, int]) -> dict:
    """Kesepakatan mentah Fase B, jumlah butir per besar selisih, dan tabel silang skor (baris = anotator 1,
    kolom = anotator 2)."""
    butir = list(s1)
    gap = Counter(abs(s1[b] - s2[b]) for b in butir)
    return {"setuju": gap[0], "dari": len(butir), "persen": 100 * gap[0] / len(butir), "selisih_1": gap[1],
            "selisih_2": gap[2],
            "tabel_silang": {a: {c: sum(s1[b] == a and s2[b] == c for b in butir) for c in SCORES} for a in SCORES}}


def method_difference(key: dict, s1: dict[str, int], s2: dict[str, int], ids: list[str], n: int = N_DRAW,
                      random_state: int = RANDOM_STATE) -> dict | None:
    """Selisih rerata skor deep agent - LLM tunggal dengan dua pembobotan: per item (rerata atas item seperti tabel per
    metode) dan per baris log (rerata atas baris dari selisih rerata skor kedua metode pada baris yang sama; hanya baris
    yang punya item dari kedua metode). IK bootstrap klaster per baris log seperti docstring modul; draw yang sama untuk
    kedua pembobotan dan ketiga skor. None bila salah satu metode tidak ada di kunci."""
    methods = ("agent", "llm_single")
    if not all(any(a["metode"] == m for x in key["butir"] for a in x["asal"]) for m in methods):
        return None
    rows = sorted(set(ids))
    pos = {r: k for k, r in enumerate(rows)}
    count = np.zeros((len(methods), len(rows)))
    total = np.zeros((len(methods), 3, len(rows)))  # metode x (rata-rata dua anotator, anotator 1, anotator 2) x baris
    for x in key["butir"]:
        scores = ((s1[x["butir"]] + s2[x["butir"]]) / 2, s1[x["butir"]], s2[x["butir"]])
        for a in x["asal"]:
            if a["metode"] in methods:
                m, r = methods.index(a["metode"]), pos[x["id"]]
                count[m, r] += 1
                total[m, :, r] += scores
    paired = (count > 0).all(axis=0)
    per_row = total[:, :, paired] / count[:, None, paired]
    row_difference = per_row[0] - per_row[1]  # skor x baris berpasangan

    def difference(weights: np.ndarray) -> np.ndarray:
        mean = (total @ weights) / (count @ weights)[:, None]
        rows_drawn = weights[paired]
        return np.concatenate([mean[0] - mean[1], row_difference @ rows_drawn / rows_drawn.sum()])

    rng = np.random.default_rng(random_state)
    draws = np.array([difference(np.bincount(rng.integers(0, len(rows), len(rows)), minlength=len(rows)))
                      for _ in range(n)])
    draws[np.abs(draws) < 1e-9] = 0.0  # selisih yang nol secara aritmetika bisa tersisa ±1e-17 (pembulatan); p memakai nol
    observed = difference(np.ones(len(rows)))
    out = {"arah": "deep agent - LLM tunggal", "klaster": len(rows), "baris_berpasangan": int(paired.sum()),
           "n_draw": n, "per_item": {}, "per_baris": {}}
    cells = [(w, s) for w in ("per_item", "per_baris") for s in ("rata_rata_dua_anotator", "anotator1", "anotator2")]
    for k, (weighting, name) in enumerate(cells):
        low, high = np.percentile(draws[:, k], [2.5, 97.5])
        p = min(1.0, 2 * min(float((draws[:, k] <= 0).mean()), float((draws[:, k] >= 0).mean())))
        out[weighting][name] = {"selisih": float(observed[k]), "ik_bawah": float(low), "ik_atas": float(high), "p": p}
    return out


def phase_b_report(key: dict, s1: dict[str, int], s2: dict[str, int], gate: dict[str, str],
                   ids: list[str] | None = None) -> dict:
    butir = list(s1)
    ids = ids if ids is not None else sorted({x["id"] for x in key["butir"]})
    out = {"butir": len(butir), "kappa_berbobot_kuadratik": float(cohen_kappa_score(
        [s1[b] for b in butir], [s2[b] for b in butir], weights="quadratic", labels=list(SCORES))),
        "kesepakatan_mentah": score_agreement(s1, s2), "run": key.get("run", {}), "metode": {}, "metode_per_label": {},
        "selisih_metode": method_difference(key, s1, s2, ids)}
    found = {a["metode"] for x in key["butir"] for a in x["asal"]}
    for method in [m for m in METHOD_RUNS if m in found] + sorted(found - set(METHOD_RUNS)):
        inst = [(x, a) for x in key["butir"] for a in x["asal"] if a["metode"] == method]
        out["metode"][method] = _relevance(inst, s1, s2, gate)
        labels = {x["label_prediksi"] for x, _ in inst}
        out["metode_per_label"][method] = {
            label: _relevance([(x, a) for x, a in inst if x["label_prediksi"] == label], s1, s2, gate)
            for label in [c for c in CLASSES if c in labels] + sorted(labels - set(CLASSES))}
    return out


def _f(x, digits: int = 3) -> str:
    return "–" if x is None else f"{x:.{digits}f}"


def _dist(d: dict) -> str:
    return "/".join(str(d[s]) for s in SCORES)


def _scores(x: dict) -> str:
    return (f"{x['item']} | {x['baris']} | {_f(x['rerata'], 2)} | {_f(x['rerata_anotator1'], 2)} | "
            f"{_f(x['rerata_anotator2'], 2)} | {_dist(x['distribusi_anotator1'])} | {_dist(x['distribusi_anotator2'])}")


def _p(p: float) -> str:
    return "<0.001" if p < 0.001 else _f(p)  # seperti src/eval/revisi.py


def _shown(path: Path) -> str:
    path = path.resolve()
    return str(path.relative_to(ROOT)) if path.is_relative_to(ROOT) else path.name


def render(a: dict, b: dict, files: dict[str, Path]) -> str:
    out = ["# Audit manusia — hasil anotasi", "",
           "Berkas: " + ", ".join(f"{k} = {_shown(v)}" for k, v in files.items()), "",
           "## Fase A — label kelas (100 baris)", ""]
    ragu = a.get("ragu", {})
    out += [f"- Kesepakatan mentah antar anotator: {_f(a['persen_setuju'], 1)}% ({a['n'] - len(a['perlu_diskusi'])} dari "
            f"{a['n']} baris); Cohen's κ antar anotator (tak berbobot): {_f(a['kappa_antar_anotator'])}.",
            f"- Anotator 1 vs label gerbang bukti: κ {_f(a['kappa_anotator1_vs_gerbang'])}, kesepakatan mentah "
            f"{_f(a['persen_setuju_anotator1_vs_gerbang'], 1)}%. Anotator 2 vs label gerbang: κ "
            f"{_f(a['kappa_anotator2_vs_gerbang'])}, kesepakatan mentah {_f(a['persen_setuju_anotator2_vs_gerbang'], 1)}%.",
            f"- Perlu diskusi (anotator tidak setuju): {len(a['perlu_diskusi'])} baris.",
            f"- Konsensus berbeda dari label gerbang: {len(a['konsensus_beda_gerbang'])} baris.",
            f"- Baris dengan catatan RAGU (kolom catatan berawalan RAGU): anotator 1 {ragu.get('anotator1', '–')}, "
            f"anotator 2 {ragu.get('anotator2', '–')}.", ""]
    if a["perlu_diskusi"]:
        out += ["Baris yang perlu diskusi (label gerbang sengaja tidak ditampilkan):", "",
                "| id | anotator 1 | anotator 2 |", "|---|---|---|",
                *[f"| {x['id']} | {x['anotator1']} | {x['anotator2']} |" for x in a["perlu_diskusi"]], ""]
    if a["konsensus_beda_gerbang"]:
        out += ["Ringkasan per kelas (gerbang -> konsensus): " + "; ".join(
            f"{pair}: {n}" for pair, n in differing_by_class(a).items()) + ". Rincian per id di "
            "results/revisi_anotasi_per_baris.json.", ""]
    matrix = a["matriks_konfusi_konsensus_vs_gerbang"]
    columns = [*CLASSES, DISCUSS]
    out += ["Matriks konfusi konsensus vs label gerbang (baris = label gerbang, kolom = konsensus kedua anotator; "
            f"\"{DISCUSS}\" = kedua anotator berbeda):", "",
            "| gerbang / konsensus | " + " | ".join(columns) + " | total |", "|---" * (len(columns) + 2) + "|",
            *[f"| {g} | " + " | ".join(str(row[c]) for c in columns) + f" | {sum(row.values())} |"
              for g, row in matrix.items()],
            "| total | " + " | ".join(str(sum(row[c] for row in matrix.values())) for c in columns)
            + f" | {sum(sum(row.values()) for row in matrix.values())} |", ""]
    agree = b.get("kesepakatan_mentah")
    out += ["## Fase B — relevansi bukti (skala 0–2)", "",
            f"- Butir unik: {b['butir']}; κ berbobot kuadratik antar anotator: {_f(b['kappa_berbobot_kuadratik'])}."]
    if agree:
        out.append(f"- Kesepakatan mentah antar anotator: {_f(agree['persen'], 1)}% ({agree['setuju']} dari "
                   f"{agree['dari']} butir); dari {agree['dari'] - agree['setuju']} butir yang berbeda, "
                   f"{agree['selisih_1']} berselisih 1 dan {agree['selisih_2']} berselisih 2.")
    out += ["- Run: " + "; ".join(f"{METHOD_NAMES.get(m, m)} = {run}" for m, run in b.get("run", {}).items()) + ".", "",
            "Per metode:", "",
            "| metode | item | baris | rerata (2 anotator) | rerata anotator 1 | rerata anotator 2 | skor 0/1/2 anotator 1 "
            "| skor 0/1/2 anotator 2 | rerata, label benar (item) | rerata, label salah (item) |",
            "|---|---|---|---|---|---|---|---|---|---|"]
    for method, x in b["metode"].items():
        out.append(f"| {METHOD_NAMES.get(method, method)} | {_scores(x)} | {_f(x['rerata_label_benar'], 2)} "
                   f"({x['item_label_benar']}) | {_f(x['rerata_label_salah'], 2)} ({x['item_label_salah']}) |")
    out += ["", "Per metode dan label prediksi:", "",
            "| metode | label_prediksi | item | baris | rerata (2 anotator) | rerata anotator 1 | rerata anotator 2 "
            "| skor 0/1/2 anotator 1 | skor 0/1/2 anotator 2 |", "|---|---|---|---|---|---|---|---|---|"]
    for method, by_label in b.get("metode_per_label", {}).items():
        out += [f"| {METHOD_NAMES.get(method, method)} | {label} | {_scores(x)} |" for label, x in by_label.items()]
    if agree:
        table = agree["tabel_silang"]
        out += ["", "Tabel silang skor per butir (baris = anotator 1, kolom = anotator 2):", "",
                "| anotator 1 / anotator 2 | " + " | ".join(map(str, SCORES)) + " | total |",
                "|---" * (len(SCORES) + 2) + "|",
                *[f"| {s} | " + " | ".join(str(table[s][t]) for t in SCORES) + f" | {sum(table[s].values())} |"
                  for s in SCORES],
                "| total | " + " | ".join(str(sum(table[s][t] for s in SCORES)) for t in SCORES)
                + f" | {agree['dari']} |"]
    diff = b.get("selisih_metode")
    if diff:
        def cell(x: dict) -> str:
            return f"{_f(x['selisih'])} | [{_f(x['ik_bawah'])}; {_f(x['ik_atas'])}] | {_p(x['p'])}"

        out += ["", f"Selisih rerata skor relevansi deep agent − LLM tunggal, IK 95% bootstrap klaster per baris log "
                f"({diff['klaster']} baris, {diff['n_draw']} draw, default_rng({RANDOM_STATE}), draw yang sama untuk "
                "semua sel; IK = persentil 2,5/97,5; p = 2 × porsi terkecil draw dengan selisih ≤ 0 atau ≥ 0). "
                "Per item = rerata atas item evidence seperti tabel per metode; per baris log = rerata atas baris dari "
                "selisih rerata skor kedua metode pada baris yang sama, setiap baris berbobot sama "
                f"({diff['baris_berpasangan']} baris memiliki item dari kedua metode):", "",
                "| skor | per item: selisih | IK 95% | p | per baris log: selisih | IK 95% | p |",
                "|---|---|---|---|---|---|---|",
                *[f"| {label} | {cell(diff['per_item'][k])} | {cell(diff['per_baris'][k])} |"
                  for k, label in (("rata_rata_dua_anotator", "rata-rata dua anotator"), ("anotator1", "anotator 1"),
                                   ("anotator2", "anotator 2"))]]
    out += ["", "Item = kemunculan item evidence pada keluaran metode; item yang identik di kedua metode dinilai sekali dan "
            "dihitung untuk keduanya, dan item yang muncul dua kali pada keluaran yang sama dihitung dua kali. Rerata "
            "(2 anotator) = rerata atas item dari skor rata-rata kedua anotator. Skor 0/1/2 = jumlah item per skor. Label "
            "benar/salah menurut label gerbang bukti."]
    return "\n".join(out) + "\n"


def differing_by_class(a: dict) -> dict[str, int]:
    counts = Counter((x["gerbang"], x["konsensus"]) for x in a["konsensus_beda_gerbang"])
    return {f"{g} -> {c}": n for (g, c), n in sorted(counts.items())}


def split_report(a: dict) -> tuple[dict, dict]:
    """Fase A untuk reports/ (tanpa label gerbang dan konsensus per id) dan untuk results/ (per id)."""
    public = {k: v for k, v in a.items() if k not in ("konsensus", "konsensus_beda_gerbang")}
    public["jumlah_konsensus_beda_gerbang"] = len(a["konsensus_beda_gerbang"])
    public["konsensus_beda_gerbang_per_kelas"] = differing_by_class(a)
    per_row = {"keterangan": "Rincian per id audit Fase A: label konsensus kedua anotator (atau 'perlu diskusi') dan "
                             "label gerbang pada baris yang konsensusnya berbeda. Ringkasan agregat di "
                             "reports/revisi_anotasi.md.",
               "konsensus": a["konsensus"], "konsensus_beda_gerbang": a["konsensus_beda_gerbang"]}
    return public, per_row


def write_outputs(a: dict, b: dict, text: str, out: Path, per_row_path: Path = PER_ROW_PATH) -> None:
    public, per_row = split_report(a)
    out.parent.mkdir(parents=True, exist_ok=True)
    out.write_text(text, encoding="utf-8")
    out.with_suffix(".json").write_text(json.dumps({"fase_a": public, "fase_b": b}, ensure_ascii=False, indent=1) + "\n",
                                        encoding="utf-8")
    per_row_path.parent.mkdir(parents=True, exist_ok=True)
    per_row_path.write_text(json.dumps(per_row, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def compute(files: dict[str, Path], ids_path: Path = IDS_PATH, key_path: Path = KEY_PATH,
            sample_path: Path = SAMPLE_PATH) -> tuple[dict, dict, str]:
    ids = json.loads(ids_path.read_text(encoding="utf-8"))["ids"]
    key = json.loads(key_path.read_text(encoding="utf-8"))
    sample = load_sample(sample_path)
    gate = dict(zip(sample["id"], sample["label"]))
    a = phase_a_report(labels_a(files["A1"], ids), labels_a(files["A2"], ids), gate, ids,
                       {"anotator1": ragu_rows(files["A1"]), "anotator2": ragu_rows(files["A2"])})
    b = phase_b_report(key, scores_b(files["B1"], key, "anotator1"), scores_b(files["B2"], key, "anotator2"), gate, ids)
    return a, b, render(a, b, files)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Paket audit manusia (Fase A label, Fase B relevansi bukti)")
    sub = parser.add_subparsers(dest="perintah", required=True)
    sub.add_parser("siapkan", help="buat daftar id, empat berkas anotasi, dan kunci Fase B")
    sub.add_parser("kunci", help="buat daftar id dan kunci Fase B saja (tanpa berkas anotasi)")
    count = sub.add_parser("hitung", help="hitung hasil dari berkas yang sudah diisi")
    for name in ("a1", "a2", "b1", "b2"):
        count.add_argument(f"--{name}", type=Path, required=True)
    count.add_argument("--out", type=Path, default=ROOT / "reports" / "revisi_anotasi.md")
    count.add_argument("--per-baris", type=Path, default=PER_ROW_PATH, help="rincian per id (label gerbang) di results/")
    args = parser.parse_args(argv)
    if args.perintah == "siapkan":
        done = prepare()
        print(f"[anotasi] {len(done['ids'])} baris, {done['butir']} butir Fase B; berkas: "
              + ", ".join(str(p.relative_to(ROOT)) for p in [*done["fase_a"].values(), *done["fase_b"].values()])
              + f"; kunci {done['kunci'].relative_to(ROOT)}")
        return 0
    if args.perintah == "kunci":
        done = prepare_key()
        print(f"[anotasi] {len(done['ids'])} baris, {done['butir']} butir Fase B; daftar id "
              f"{IDS_PATH.relative_to(ROOT)}, kunci {done['kunci'].relative_to(ROOT)}")
        return 0
    files = {"A1": args.a1, "A2": args.a2, "B1": args.b1, "B2": args.b2}
    a, b, text = compute(files)
    print(text)
    write_outputs(a, b, text, args.out, args.per_baris)
    return 0


if __name__ == "__main__":
    sys.exit(main())
